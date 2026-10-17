# A folder of files with an access file: each supported type is turned into text, the last
# matching access line gives a file its readers, a file with no line or a link stops the load
# before anything is written, and a file load deletes only files, never documents.jsonl entries.

import json
import zipfile
from collections.abc import Iterator
from pathlib import Path

import openpyxl
import psycopg
import pytest
from pptx import Presentation
from pptx.util import Inches
from psycopg.rows import TupleRow

from needtoknow.auth import EVERYONE
from needtoknow.company import ID_PREFIX, load_documents, read_documents
from needtoknow.files import FILE_PREFIX, FilesError, read_files
from needtoknow.retrieval import search

Connection = psycopg.Connection[TupleRow]
MAX_BYTES = 1024 * 1024
DOCX_TYPES = (
    '<?xml version="1.0" encoding="UTF-8"?>'
    '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
    '<Default Extension="rels" '
    'ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
    '<Default Extension="xml" ContentType="application/xml"/>'
    '<Override PartName="/word/document.xml" ContentType="application/'
    'vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/>'
    "</Types>"
)
DOCX_RELS = (
    '<?xml version="1.0" encoding="UTF-8"?>'
    '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
    '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/'
    'relationships/officeDocument" Target="word/document.xml"/>'
    "</Relationships>"
)
DOCX_BODY = (
    '<?xml version="1.0" encoding="UTF-8"?>'
    '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
    "<w:body><w:p><w:r><w:t>{text}</w:t></w:r></w:p></w:body></w:document>"
)


@pytest.fixture
def owner_clean(owner: Connection) -> Iterator[Connection]:
    yield owner
    with owner.transaction():
        for prefix in (FILE_PREFIX, ID_PREFIX):
            owner.execute("DELETE FROM chunks WHERE starts_with(doc_id, %s)", [prefix])
            owner.execute("DELETE FROM doc_access WHERE starts_with(doc_id, %s)", [prefix])
            owner.execute("DELETE FROM documents WHERE starts_with(id, %s)", [prefix])


def _access(path: Path, *entries: dict[str, object]) -> Path:
    path.write_text("".join(json.dumps(entry) + "\n" for entry in entries), encoding="utf-8")
    return path


def _docx(path: Path, text: str) -> None:
    with zipfile.ZipFile(path, "w") as archive:
        archive.writestr("[Content_Types].xml", DOCX_TYPES)
        archive.writestr("_rels/.rels", DOCX_RELS)
        archive.writestr("word/document.xml", DOCX_BODY.format(text=text))


def _xlsx(path: Path) -> None:
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.title = "Budget"
    sheet.append(["Team", "Euros"])
    sheet.append(["Platform", 120000])
    workbook.save(path)


def _pptx(path: Path, title: str, text: str) -> None:
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[5])
    slide.shapes.title.text = title
    box = slide.shapes.add_textbox(Inches(1), Inches(2), Inches(6), Inches(1))
    box.text_frame.text = text
    deck.save(str(path))


def _ids(owner: Connection, prefix: str) -> list[str]:
    rows = owner.execute(
        "SELECT id FROM documents WHERE starts_with(id, %s) ORDER BY id", [prefix]
    ).fetchall()
    return [row[0] for row in rows]


def test_each_supported_type_is_read_as_text(tmp_path: Path) -> None:
    folder = tmp_path / "files"
    (folder / "hr").mkdir(parents=True)
    _docx(folder / "hr" / "leave-policy.docx", "Employees take 25 days of paid leave.")
    _xlsx(folder / "budget.xlsx")
    _pptx(folder / "kickoff.pptx", "Project kickoff", "The launch is set for May.")
    (folder / "office.html").write_text(
        "<html><body><h1>Lisbon office</h1><p>Doors open at eight.</p></body></html>",
        encoding="utf-8",
    )
    (folder / "notes.md").write_text("# Parking\n\nParking is on level two.\n", encoding="utf-8")
    (folder / "wifi.txt").write_text("The guest network is calder-guest.\n", encoding="utf-8")
    access = _access(tmp_path / "access.jsonl", {"path": "*", "public": True})

    found = read_files(folder, access, MAX_BYTES)

    documents = {document.id: document for document in found.documents}
    assert sorted(documents) == [
        "file-budget.xlsx",
        "file-hr/leave-policy.docx",
        "file-kickoff.pptx",
        "file-notes.md",
        "file-office.html",
        "file-wifi.txt",
    ]
    assert documents["file-hr/leave-policy.docx"].title == "leave-policy"
    assert "25 days of paid leave" in documents["file-hr/leave-policy.docx"].body
    assert "| Platform | 120000 |" in documents["file-budget.xlsx"].body
    assert documents["file-kickoff.pptx"].title == "Project kickoff"
    assert "The launch is set for May." in documents["file-kickoff.pptx"].body
    assert documents["file-office.html"].title == "Lisbon office"
    assert "Doors open at eight." in documents["file-office.html"].body
    assert documents["file-notes.md"].title == "Parking"
    assert documents["file-wifi.txt"].title == "wifi"
    assert documents["file-wifi.txt"].body == "The guest network is calder-guest.\n"


def test_the_last_matching_access_line_gives_the_readers(tmp_path: Path) -> None:
    folder = tmp_path / "files"
    (folder / "finance" / "2026").mkdir(parents=True)
    (folder / "finance" / "2026" / "q1.md").write_text("Revenue grew.", encoding="utf-8")
    (folder / "finance" / "open.md").write_text("The finance desk is on floor 3.", encoding="utf-8")
    (folder / "welcome.txt").write_text("Welcome.", encoding="utf-8")
    access = _access(
        tmp_path / "access.jsonl",
        {"path": "*", "public": True},
        {"path": "finance/*", "groups": ["finance"]},
        {"path": "finance/open.md", "groups": ["finance"], "users": ["alice"]},
    )

    found = read_files(folder, access, MAX_BYTES)

    readers = {document.id: document.readers for document in found.documents}
    assert readers == {
        "file-finance/2026/q1.md": frozenset({"group:finance"}),
        "file-finance/open.md": frozenset({"group:finance", "user:alice"}),
        "file-welcome.txt": frozenset({EVERYONE}),
    }


def test_other_types_and_files_over_the_cap_are_skipped_and_counted(tmp_path: Path) -> None:
    folder = tmp_path / "files"
    folder.mkdir()
    (folder / "logo.png").write_bytes(b"\x89PNG\r\n\x1a\n")
    (folder / "archive.zip").write_bytes(b"PK\x05\x06" + bytes(18))
    (folder / "large.txt").write_text("x" * 101, encoding="utf-8")
    (folder / "small.txt").write_text("Small.", encoding="utf-8")
    access = _access(tmp_path / "access.jsonl", {"path": "*", "public": True})

    found = read_files(folder, access, 100)

    assert [document.id for document in found.documents] == ["file-small.txt"]
    assert (found.unsupported, found.oversized) == (2, 1)


def test_a_file_no_access_line_matches_stops_the_load_before_anything_is_written(
    owner_clean: Connection, tmp_path: Path
) -> None:
    folder = tmp_path / "files"
    (folder / "hr").mkdir(parents=True)
    (folder / "hr" / "salaries.md").write_text("Salaries are reviewed in June.", encoding="utf-8")
    (folder / "welcome.txt").write_text("Welcome.", encoding="utf-8")
    access = _access(tmp_path / "access.jsonl", {"path": "*.txt", "public": True})

    with pytest.raises(FilesError, match=r"gives no readers for hr/salaries\.md$"):
        load_documents(
            owner_clean, read_files(folder, access, MAX_BYTES).documents, FILE_PREFIX, True
        )

    assert _ids(owner_clean, FILE_PREFIX) == []


def test_a_symbolic_link_is_refused(tmp_path: Path) -> None:
    secret = tmp_path / "secret.txt"
    secret.write_text("Outside the folder.", encoding="utf-8")
    folder = tmp_path / "files"
    folder.mkdir()
    (folder / "linked.txt").symlink_to(secret)
    access = _access(tmp_path / "access.jsonl", {"path": "*", "public": True})

    with pytest.raises(FilesError, match="linked.txt is a symbolic link"):
        read_files(folder, access, MAX_BYTES)


def test_a_linked_folder_is_refused(tmp_path: Path) -> None:
    outside = tmp_path / "outside"
    outside.mkdir()
    (outside / "secret.txt").write_text("Outside the folder.", encoding="utf-8")
    folder = tmp_path / "files"
    folder.mkdir()
    (folder / "shared").symlink_to(outside, target_is_directory=True)
    access = _access(tmp_path / "access.jsonl", {"path": "*", "public": True})

    with pytest.raises(FilesError, match="shared is a symbolic link"):
        read_files(folder, access, MAX_BYTES)


@pytest.mark.parametrize("pattern", ["/etc/*", "../*", "hr/../../*"])
def test_an_access_path_outside_the_folder_is_refused(tmp_path: Path, pattern: str) -> None:
    folder = tmp_path / "files"
    folder.mkdir()
    access = _access(tmp_path / "access.jsonl", {"path": pattern, "public": True})

    with pytest.raises(FilesError, match="line 1: path .* must stay inside the folder"):
        read_files(folder, access, MAX_BYTES)


def test_a_file_is_found_only_by_its_readers(
    owner_clean: Connection, app: Connection, tmp_path: Path
) -> None:
    folder = tmp_path / "files"
    folder.mkdir()
    (folder / "move.md").write_text(
        "# Lisbon office move\n\nThe Lisbon office moves to Rua Augusta in March.\n",
        encoding="utf-8",
    )
    access = _access(tmp_path / "access.jsonl", {"path": "*", "users": ["files-test-alice"]})
    load_documents(owner_clean, read_files(folder, access, MAX_BYTES).documents, FILE_PREFIX, False)

    question = "When does the Lisbon office move?"
    alice = search(app, ["user:files-test-alice"], question, 5)
    bob = search(app, ["user:files-test-bob"], question, 5)

    assert [hit.doc_id for hit in alice] == ["file-move.md"]
    assert bob == []


def test_deleting_missing_files_keeps_documents_from_the_jsonl_file(
    owner_clean: Connection, tmp_path: Path
) -> None:
    folder = tmp_path / "files"
    folder.mkdir()
    (folder / "parking.md").write_text("Parking is on level two.", encoding="utf-8")
    (folder / "wifi.txt").write_text("The guest network is calder-guest.", encoding="utf-8")
    access = _access(tmp_path / "access.jsonl", {"path": "*", "public": True})
    jsonl = tmp_path / "documents.jsonl"
    jsonl.write_text(
        json.dumps({"id": "office-move", "title": "Move", "body": "In March.", "public": True})
        + "\n",
        encoding="utf-8",
    )
    load_documents(owner_clean, read_documents(jsonl), ID_PREFIX, False)
    load_documents(owner_clean, read_files(folder, access, MAX_BYTES).documents, FILE_PREFIX, False)

    (folder / "wifi.txt").unlink()
    files = load_documents(
        owner_clean, read_files(folder, access, MAX_BYTES).documents, FILE_PREFIX, True
    )
    jsonl.write_text("", encoding="utf-8")
    company = load_documents(owner_clean, read_documents(jsonl), ID_PREFIX, True)

    assert (files.deleted, company.deleted) == (1, 1)
    assert _ids(owner_clean, FILE_PREFIX) == ["file-parking.md"]
    assert _ids(owner_clean, ID_PREFIX) == []
